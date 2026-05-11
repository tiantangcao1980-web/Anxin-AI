# -*- coding: utf-8 -*-
"""xlsx 读取工具 - 返回 pandas DataFrame，便于后续 analyze。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def read_sheet(
    path: str | Path,
    sheet_name: str | int | None = None,
    *,
    header: int | None = 0,
) -> pd.DataFrame:
    """读取单个 sheet 为 DataFrame。

    Args:
        path: xlsx 文件路径
        sheet_name: sheet 名或索引；None 取第一个
        header: 表头行（行号 0-indexed），None 表示无表头
    """
    if sheet_name is None:
        sheet_name = 0
    return pd.read_excel(path, sheet_name=sheet_name, header=header, engine="openpyxl")


def read_all_sheets(
    path: str | Path,
    *,
    header: int | None = 0,
) -> dict[str, pd.DataFrame]:
    """读取所有 sheet，返回 dict[sheet_name, DataFrame]。"""
    return pd.read_excel(path, sheet_name=None, header=header, engine="openpyxl")


def read_formulas(
    path: str | Path,
    sheet_name: str | None = None,
) -> dict[str, str]:
    """提取所有公式，返回 {cell_coord: formula}。

    Args:
        path: xlsx 路径
        sheet_name: 指定 sheet；None 表示遍历所有 sheet
    """
    wb = load_workbook(filename=str(path), data_only=False)
    result: dict[str, str] = {}

    sheets = [sheet_name] if sheet_name else wb.sheetnames
    for name in sheets:
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    key = f"{name}!{cell.coordinate}" if not sheet_name else cell.coordinate
                    result[key] = v
    return result


def read_charts_info(path: str | Path) -> list[dict]:
    """提取所有图表的元信息（类型、所在 sheet、anchor）。"""
    wb = load_workbook(filename=str(path))
    out: list[dict] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for chart in getattr(ws, "_charts", []):
            out.append(
                {
                    "sheet": sheet_name,
                    "type": type(chart).__name__,
                    "title": getattr(chart, "title", None) and str(chart.title),
                    "anchor": str(getattr(chart, "anchor", "")),
                }
            )
    return out
