# -*- coding: utf-8 -*-
"""xlsx 编辑工具 - 单元格更新、行插入、公式重算触发。"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .create import _check_formula_safe


def update_cell(
    path: str | Path,
    sheet: str,
    cell: str,
    value: Any,
) -> None:
    """更新单个单元格并保存。

    若 value 是公式（= 起头），走安全检查。
    """
    if isinstance(value, str) and value.startswith("="):
        _check_formula_safe(value)
    wb = load_workbook(filename=str(path))
    ws = wb[sheet]
    ws[cell] = value
    wb.save(str(path))


def insert_rows(
    path: str | Path,
    sheet: str,
    position: int,
    rows: Iterable[Iterable[Any]],
) -> None:
    """在 position 之前插入若干行。

    Args:
        path: xlsx 路径
        sheet: sheet 名
        position: 1-indexed 行号；新行会插在这一行之前
        rows: 二维数据
    """
    wb = load_workbook(filename=str(path))
    ws = wb[sheet]
    rows_list = [list(r) for r in rows]
    if not rows_list:
        return
    ws.insert_rows(position, amount=len(rows_list))
    for r_offset, row_data in enumerate(rows_list):
        for c_offset, value in enumerate(row_data, start=1):
            if isinstance(value, str) and value.startswith("="):
                _check_formula_safe(value)
            ws.cell(row=position + r_offset, column=c_offset, value=value)
    wb.save(str(path))


def recalc_formulas(path: str | Path) -> None:
    """触发 Excel 下次打开时重新计算所有公式。

    openpyxl 不会自己求值，但可以通过清空 cached value 让 Excel 打开时重算。
    实现方式：以 data_only=False 读，再原样存。
    """
    wb = load_workbook(filename=str(path), data_only=False)
    # 标记 calc properties full_calc_on_load
    if wb.calculation is not None:
        wb.calculation.fullCalcOnLoad = True
    wb.save(str(path))
