# -*- coding: utf-8 -*-
"""tests/test_edit.py - 单元格更新、行插入、公式保留。"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from skills.office.xlsx.helpers import create, edit


def _make(path):
    create.create_xlsx(
        path,
        {
            "S": [
                ["A", "B"],
                [1, 2],
                [3, 4],
                ["合计", "=SUM(B2:B3)"],
            ]
        },
    )


def test_edit_preserves_formulas(tmp_path):
    """更新非公式单元格后，已有公式仍在。"""
    out = tmp_path / "e.xlsx"
    _make(out)

    edit.update_cell(out, "S", "A1", "字段")

    wb = load_workbook(out, data_only=False)
    ws = wb["S"]
    assert ws["A1"].value == "字段"
    assert ws["B4"].value == "=SUM(B2:B3)"  # 公式没丢


def test_insert_rows_shifts_data(tmp_path):
    out = tmp_path / "e.xlsx"
    _make(out)

    edit.insert_rows(out, "S", position=2, rows=[[0, 0]])

    wb = load_workbook(out, data_only=False)
    ws = wb["S"]
    # 新行在第 2 行
    assert ws["A2"].value == 0
    # 原 [1,2] 行被推到第 3 行
    assert ws["A3"].value == 1
    assert ws["B3"].value == 2


def test_update_cell_rejects_dangerous_formula(tmp_path):
    out = tmp_path / "e.xlsx"
    _make(out)
    with pytest.raises(ValueError, match="formula injection rejected"):
        edit.update_cell(out, "S", "C1", "=cmd|/c calc.exe!A1")
