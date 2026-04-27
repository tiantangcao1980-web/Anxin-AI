# -*- coding: utf-8 -*-
"""tests/test_styling.py - 表头样式 + 条件格式 + 列宽。"""

from __future__ import annotations

from openpyxl import Workbook, load_workbook

from skills.office.xlsx.helpers import create, styling


def test_styling_header(tmp_path):
    """apply_header_style 后，第 1 行应该 bold + 灰底。"""
    out = tmp_path / "h.xlsx"
    wb = create.create_xlsx(
        out,
        {"S": [["A", "B", "C"], [1, 2, 3]]},
    )
    ws = wb["S"]
    styling.apply_header_style(ws, row=1)
    wb.save(out)

    wb2 = load_workbook(out)
    ws2 = wb2["S"]
    a1 = ws2["A1"]
    assert a1.font.bold is True
    # 颜色比对：openpyxl 把 fgColor 包成 Color 对象
    assert a1.fill.fgColor.rgb in {"00D9D9D9", "D9D9D9", "FFD9D9D9"}


def test_apply_conditional_format_high(tmp_path):
    out = tmp_path / "c.xlsx"
    wb = create.create_xlsx(out, {"S": [["v"], [1], [10], [100]]})
    ws = wb["S"]
    styling.apply_conditional_format(ws, "A2:A4", "high", threshold=5)
    wb.save(out)

    wb2 = load_workbook(out)
    ws2 = wb2["S"]
    # 至少有一条规则
    rules = list(ws2.conditional_formatting._cf_rules.values())
    assert len(rules) >= 1


def test_auto_fit_columns_sets_width(tmp_path):
    out = tmp_path / "w.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "短"
    ws["A2"] = "这是一段比较长的文本用来测试自适应列宽"
    styling.auto_fit_columns(ws)
    # width 应该被设置过且不为 None
    assert ws.column_dimensions["A"].width is not None
    assert ws.column_dimensions["A"].width > 5
